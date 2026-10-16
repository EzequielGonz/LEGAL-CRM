import os, datetime, csv, textwrap
from PIL import Image, ImageDraw, ImageFont
from contenido import RUBROS, PIEZAS
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
F="/usr/share/fonts/opentype/inter/"
def font(w,s): return ImageFont.truetype(F+{"b":"Inter-Bold.otf","sb":"Inter-SemiBold.otf","r":"Inter-Regular.otf","m":"Inter-Medium.otf"}[w], s)
def hexrgb(h): return tuple(int(h[i:i+2],16) for i in (1,3,5))
def lum(c):
    def ch(v):
        v/=255; return v/12.92 if v<=0.03928 else ((v+0.055)/1.055)**2.4
    r,g,b=[ch(x) for x in c]; return 0.2126*r+0.7152*g+0.0722*b
def contrast(a,b):
    la,lb=sorted([lum(a),lum(b)],reverse=True); return (la+0.05)/(lb+0.05)
W,H=1080,1350
def wrap(d,text,f,maxw):
    out=[]
    for para in text.split("\n"):
        words=para.split(); line=""
        for w in words:
            t=(line+" "+w).strip()
            if d.textlength(t,font=f)<=maxw: line=t
            else: out.append(line); line=w
        out.append(line)
    return out
def pieza(rubro, i, pilar, tit, baj, path):
    oscuro,acento,claro=[hexrgb(c) for c in RUBROS[rubro]["colores"]]
    dark = i%2==0
    bg,fg=(oscuro,claro) if dark else (claro,oscuro)
    img=Image.new("RGB",(W,H),bg); d=ImageDraw.Draw(img)
    # forma decorativa
    d.ellipse([W-380,H-600,W+140,H-80],fill=acento if dark else tuple(int(a*0.25+b*0.75) for a,b in zip(acento,claro)))
    d.rectangle([96,180,96+120,180+12],fill=acento)
    d.text((96,220),pilar.upper(),font=font("sb",30),fill=acento if contrast(acento,bg)>=3 else fg)
    ft=font("b",88); y=330
    for ln in wrap(d,tit,ft,W-192):
        d.text((96,y),ln,font=ft,fill=fg); y+=104
    y+=30; fb=font("m",44)
    for ln in wrap(d,baj,fb,W-192):
        d.text((96,y),ln,font=fb,fill=fg); y+=58
    ficticio="ficticio" in tit.lower()
    if ficticio:
        d.rounded_rectangle([96,H-300,96+420,H-240],radius=30,outline=fg,width=3)
        d.text((126,H-290),"Ejemplo ilustrativo",font=font("sb",30),fill=fg)
    barra=claro if dark else oscuro
    d.rectangle([0,H-140,W,H],fill=barra)
    d.rectangle([0,H-148,W,H-140],fill=acento)
    tc=oscuro if dark else claro
    acento=barra
    d.text((96,H-98),"Consultanos por mensaje  →",font=font("sb",40),fill=tc)
    img.save(path, optimize=True)
    return contrast(fg,bg), contrast(tc,barra)
inicio=datetime.date(2026,10,19)
dias=[]
dd=inicio
while len(dias)<12:
    if dd.weekday() in (0,2,4): dias.append(dd)
    dd+=datetime.timedelta(days=1)
os.makedirs("piezas",exist_ok=True)
wb=openpyxl.Workbook(); ws0=wb.active; ws0.title="Estrategia"
ws0.append(["Rubro","Objetivo","Público","Pilares","Tono","Formato y frecuencia","Destinos","CTA","Estado"])
for r,v in RUBROS.items():
    ws0.append([v["titulo"],v["objetivo"],v["publico"],", ".join(v["pilares"]),v["tono"],
               "Imagen 1080×1350 · lun/mié/vie 11:00 (AR) · 12 piezas en 30 días; Reels a partir del mes 2 con material real",
               "Instagram + Facebook del perfil (texto adaptado por red)","Pendiente: WhatsApp/web real del perfil","Borrador: falta identidad, logo y CTA"])
ws=wb.create_sheet("Calendario")
hdr=["ID","Rubro","Fecha","Hora (AR)","Red","Pilar","Titular de la pieza","Texto","Texto alternativo","Archivo","Estado"]
ws.append(hdr)
rows=[]; minc=99
for r,lista in PIEZAS.items():
    for i,(pilar,tit,baj,ig,fb) in enumerate(lista):
        pid=f"{r[:3]}-{i+1:02d}"
        fn=f"piezas/{r}_{i+1:02d}.png"
        c1,c2=pieza(r,i,pilar,tit,baj,fn); minc=min(minc,c1,c2)
        alt=f"Placa con fondo de color y el texto: {tit.replace(chr(10),' ')}. {baj}."
        cta="\n\n👉 Escribinos por WhatsApp: [CTA del perfil]"
        for red,txt in (("Instagram",ig+cta.replace("Escribinos por WhatsApp","Escribinos por WhatsApp (link en la bio)")+"\n\n"+RUBROS[r]["hashtags"]),("Facebook",fb+cta)):
            row=[pid,RUBROS[r]["titulo"],dias[i].isoformat(),"11:00",red,pilar,tit.replace("\n"," "),txt,alt,fn,"Borrador (falta identidad/CTA)"]
            ws.append(row); rows.append(row)
for w,col in zip([10,16,12,9,11,22,34,70,40,30,26],"ABCDEFGHIJK"): ws.column_dimensions[col].width=w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
for s in (ws0,ws):
    for c in s[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F2A44")
for col,w in zip("ABCDEFGHI",[18,40,40,40,40,40,30,30,30]): ws0.column_dimensions[col].width=w
for row in ws0.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
wb.save("calendario_redes_30_dias.xlsx")
print("piezas",len(rows)//2,"filas",len(rows),"contraste min",round(minc,2))
